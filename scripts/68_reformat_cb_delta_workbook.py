from __future__ import annotations

import argparse
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


DEFAULT_INPUT = Path(r"C:\Users\lenovo\Desktop\转债Delta分析图表数据_20260611.xlsx")
DEFAULT_OUTPUT = Path(r"C:\Users\lenovo\Desktop\转债Delta分析图表数据_20260611_可读版.xlsx")

PRICE_GROUP_ORDER = ["130以下", "130-150", "150以上"]


def _delta_col(df: pd.DataFrame) -> str:
    for col in ["实际_理论Delta均值", "实际-理论Delta均值"]:
        if col in df.columns:
            return col
    raise KeyError("Missing delta spread column.")


def build_industry_tables(path: Path) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    df = pd.read_excel(path, sheet_name="04_行业最新分化图数据")
    diff_col = _delta_col(df)
    df = df.rename(columns={diff_col: "实际-理论Delta均值"})
    numeric_cols = ["周频样本数", "实际Delta大于理论占比", "理论Delta均值", "实际Delta均值", "实际-理论Delta均值"]
    for col in numeric_cols:
        df[col] = pd.to_numeric(df[col], errors="coerce")

    latest_month = int(df["月份"].max())
    latest = (
        df[df["月份"].eq(latest_month)]
        .sort_values(["实际-理论Delta均值", "实际Delta大于理论占比", "周频样本数"], ascending=[False, False, False])
        .drop_duplicates(subset=["月份", "行业"], keep="first")
        .copy()
    )
    strong = (
        latest[latest["实际-理论Delta均值"].gt(0)]
        .sort_values(["实际-理论Delta均值", "实际Delta大于理论占比", "周频样本数"], ascending=[False, False, False])
        .head(10)
        .copy()
    )
    strong.insert(0, "排序", range(1, len(strong) + 1))
    strong.insert(1, "结论", "实际强于理论")
    strong["排序依据"] = "偏离降序；占比降序；样本数降序"

    weak = (
        latest[latest["实际-理论Delta均值"].lt(0)]
        .sort_values(["实际-理论Delta均值", "实际Delta大于理论占比", "周频样本数"], ascending=[True, True, False])
        .head(10)
        .copy()
    )
    weak.insert(0, "排序", range(1, len(weak) + 1))
    weak.insert(1, "结论", "实际弱于理论")
    weak["排序依据"] = "偏离升序；占比升序；样本数降序"

    cols = [
        "排序",
        "结论",
        "月份",
        "行业",
        "周频样本数",
        "实际Delta大于理论占比",
        "理论Delta均值",
        "实际Delta均值",
        "实际-理论Delta均值",
        "排序依据",
    ]
    return latest[cols[2:-1]].sort_values("实际-理论Delta均值", ascending=False), strong[cols], weak[cols]


def build_price_group_tables(path: Path) -> tuple[pd.DataFrame, pd.DataFrame]:
    raw = pd.read_excel(path, sheet_name="03_价格组三指标图数据")
    long = raw.melt(id_vars=["指标", "月份"], value_vars=PRICE_GROUP_ORDER, var_name="价格组", value_name="数值")
    detail = long.pivot_table(index=["月份", "价格组"], columns="指标", values="数值", aggfunc="first").reset_index()
    detail.columns.name = None
    if "实际_理论Delta均值" in detail.columns:
        detail = detail.rename(columns={"实际_理论Delta均值": "实际-理论Delta均值"})
    for col in ["实际Delta大于理论占比", "理论Delta均值", "实际Delta均值", "实际-理论Delta均值"]:
        detail[col] = pd.to_numeric(detail[col], errors="coerce")
    detail["强弱排名"] = detail.groupby("月份")["实际-理论Delta均值"].rank(method="first", ascending=False).astype(int)
    detail = detail.sort_values(["月份", "强弱排名"]).reset_index(drop=True)

    rows = []
    for month, sub in detail.groupby("月份"):
        ordered = sub.sort_values("实际-理论Delta均值", ascending=False)
        strong = ordered.iloc[0]
        weak = ordered.iloc[-1]
        rows.append(
            {
                "月份": month,
                "最强价格组": strong["价格组"],
                "最强偏离": strong["实际-理论Delta均值"],
                "最强占比": strong["实际Delta大于理论占比"],
                "最弱价格组": weak["价格组"],
                "最弱偏离": weak["实际-理论Delta均值"],
                "最弱占比": weak["实际Delta大于理论占比"],
                "强弱差": strong["实际-理论Delta均值"] - weak["实际-理论Delta均值"],
                "结论": f"{strong['价格组']}相对最强，{weak['价格组']}相对最弱",
            }
        )
    summary = pd.DataFrame(rows)
    return summary, detail


def build_summary(industry_strong: pd.DataFrame, industry_weak: pd.DataFrame, price_summary: pd.DataFrame) -> pd.DataFrame:
    rows = [
        {"项目": "行业强弱排序", "结论": "强行业按实际-理论Delta均值降序，弱行业按实际-理论Delta均值升序；占比只作为辅助确认。"},
        {"项目": "占比解释", "结论": "实际Delta大于理论占比越高，说明更多样本实际Delta超过理论；强侧应高，弱侧应低。"},
        {"项目": "最新强行业", "结论": "、".join(industry_strong["行业"].head(3).astype(str).tolist())},
        {"项目": "最新弱行业", "结论": "、".join(industry_weak["行业"].head(3).astype(str).tolist())},
        {"项目": "价格组月度表", "结论": "改为每月最强/最弱价格组摘要，避免指标纵向堆叠导致不可读。"},
    ]
    if not price_summary.empty:
        latest = price_summary.sort_values("月份").iloc[-1]
        rows.append({"项目": "最新月份价格组结论", "结论": str(latest["结论"])})
    return pd.DataFrame(rows)


def write_workbook(
    output: Path,
    summary: pd.DataFrame,
    latest_industry: pd.DataFrame,
    strong: pd.DataFrame,
    weak: pd.DataFrame,
    price_summary: pd.DataFrame,
    price_detail: pd.DataFrame,
) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="00_结论摘要", index=False)
        strong.to_excel(writer, sheet_name="01_行业强于理论Top", index=False)
        weak.to_excel(writer, sheet_name="02_行业弱于理论Top", index=False)
        latest_industry.to_excel(writer, sheet_name="03_行业全量排序", index=False)
        price_summary.to_excel(writer, sheet_name="04_价格组月度结论", index=False)
        price_detail.to_excel(writer, sheet_name="05_价格组明细排序", index=False)

    wb = load_workbook(output)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    positive_fill = PatternFill("solid", fgColor="E2F0D9")
    negative_fill = PatternFill("solid", fgColor="FCE4D6")

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.sheet_view.showGridLines = False
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.auto_filter.ref = ws.dimensions
        max_row = ws.max_row
        max_col = ws.max_column
        if max_row >= 2 and max_col >= 1:
            table_ref = f"A1:{ws.cell(max_row, max_col).coordinate}"
            table = Table(displayName=f"T_{ws.title[:20].replace('-', '_')}", ref=table_ref)
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
            ws.add_table(table)
        for col in ws.columns:
            values = [str(c.value) if c.value is not None else "" for c in col]
            width = min(max(max(len(v) for v in values) + 2, 10), 34)
            ws.column_dimensions[col[0].column_letter].width = width
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                header = ws.cell(1, cell.column).value
                if header and "占比" in str(header):
                    cell.number_format = "0.00%"
                elif header and ("Delta" in str(header) or "偏离" in str(header) or "强弱差" in str(header)):
                    cell.number_format = "0.000"
                    if isinstance(cell.value, (int, float)):
                        if cell.value > 0:
                            cell.fill = positive_fill
                        elif cell.value < 0:
                            cell.fill = negative_fill
        for col_idx in range(1, max_col + 1):
            header = str(ws.cell(1, col_idx).value)
            if "偏离" in header or "强弱差" in header:
                col_letter = ws.cell(1, col_idx).column_letter
                ws.conditional_formatting.add(
                    f"{col_letter}2:{col_letter}{max_row}",
                    ColorScaleRule(start_type="min", start_color="F8696B", mid_type="num", mid_value=0, mid_color="FFFFFF", end_type="max", end_color="63BE7B"),
                )

    if "04_价格组月度结论" in wb.sheetnames:
        ws = wb["04_价格组月度结论"]
        chart = BarChart()
        chart.type = "bar"
        chart.style = 10
        chart.title = "每月价格组强弱差"
        chart.y_axis.title = "月份"
        chart.x_axis.title = "强弱差"
        data = Reference(ws, min_col=8, min_row=1, max_row=ws.max_row)
        cats = Reference(ws, min_col=1, min_row=2, max_row=ws.max_row)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        chart.height = 7
        chart.width = 12
        ws.add_chart(chart, "K2")

    wb.save(output)


def main() -> None:
    parser = argparse.ArgumentParser(description="Reformat convertible bond Delta workbook into readable conclusion tables.")
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()

    latest_industry, strong, weak = build_industry_tables(args.input)
    price_summary, price_detail = build_price_group_tables(args.input)
    summary = build_summary(strong, weak, price_summary)
    output = args.output
    try:
        write_workbook(output, summary, latest_industry, strong, weak, price_summary, price_detail)
    except PermissionError:
        output = output.with_name(f"{output.stem}_v2{output.suffix}")
        write_workbook(output, summary, latest_industry, strong, weak, price_summary, price_detail)
    print({"output": str(output), "sheets": 6, "strong_rows": len(strong), "weak_rows": len(weak)})


if __name__ == "__main__":
    main()
