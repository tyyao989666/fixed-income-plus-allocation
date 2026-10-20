"""Build a concise Excel research/data workbook for the v43 formal release."""

from __future__ import annotations

import json
from copy import copy
from pathlib import Path

import pandas as pd
import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output/reports"
VERSION = "hierarchical_asset_itransformer_v43_clean_calendar"
OUTPUT = OUT / "v43_three_layer_framework_and_data_workbook.xlsx"

NAVY = "17365D"
BLUE = "1F4E78"
TEAL = "0F766E"
GREEN = "2E7D32"
AMBER = "B45309"
RED = "B91C1C"
PALE_BLUE = "DCE6F1"
PALE_GREEN = "E2F0D9"
PALE_AMBER = "FCE4D6"
WHITE = "FFFFFF"
THIN_GRAY = Side(style="thin", color="D9E2F3")


def read_yaml(relative: str) -> dict:
    with (ROOT / relative).open("r", encoding="utf-8") as handle:
        return yaml.safe_load(handle)


def make_sheet(workbook: Workbook, name: str, title: str, subtitle: str, last_col: str = "H") -> object:
    sheet = workbook.create_sheet(name)
    sheet.sheet_view.showGridLines = False
    sheet.merge_cells(f"A1:{last_col}1")
    sheet["A1"] = title
    sheet["A1"].font = Font(name="Microsoft YaHei", size=16, bold=True, color=WHITE)
    sheet["A1"].fill = PatternFill("solid", fgColor=NAVY)
    sheet["A1"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 28
    sheet.merge_cells(f"A2:{last_col}2")
    sheet["A2"] = subtitle
    sheet["A2"].font = Font(name="Microsoft YaHei", size=10, color="475569")
    sheet["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[2].height = 30
    return sheet


def write_table(sheet, start_row: int, headers: list[str], rows: list[list[object]], table_name: str) -> tuple[int, int]:
    for col, value in enumerate(headers, 1):
        cell = sheet.cell(start_row, col, value)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(name="Microsoft YaHei", bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row_idx, values in enumerate(rows, start_row + 1):
        for col_idx, value in enumerate(values, 1):
            cell = sheet.cell(row_idx, col_idx, value)
            cell.font = Font(name="Microsoft YaHei", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=THIN_GRAY)
    end_row = start_row + len(rows)
    end_col = len(headers)
    if rows:
        ref = f"A{start_row}:{chr(64 + end_col)}{end_row}"
        table = Table(displayName=table_name, ref=ref)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
        sheet.add_table(table)
    sheet.freeze_panes = f"A{start_row + 1}"
    sheet.auto_filter.ref = f"A{start_row}:{chr(64 + end_col)}{end_row}"
    return end_row, end_col


def format_widths(sheet, widths: dict[str, float]) -> None:
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width


def asset_data() -> pd.DataFrame:
    universe = pd.read_parquet(ROOT / "data_processed/asset_universe.parquet")
    returns = pd.read_parquet(ROOT / "data_processed/asset_returns_daily.parquet")
    returns["date"] = pd.to_datetime(returns["date"], errors="coerce")
    returns["return_1d"] = pd.to_numeric(returns["return_1d"], errors="coerce")
    coverage = returns.groupby("asset_code", as_index=False).agg(
        actual_start=("date", "min"), actual_end=("date", "max"), observations=("date", "size"),
        missing_rate=("return_1d", lambda value: float(value.isna().mean())),
    )
    out = universe.merge(coverage, on="asset_code", how="left")
    out["proxy_type"] = "直接价格/基金价格"
    out.loc[out["source_type"].astype(str).str.contains("underlying_index|wealth_index"), "proxy_type"] = "指数或财富指数代理"
    out.loc[out["source_type"].astype(str).str.contains("curve_proxy"), "proxy_type"] = "收益率曲线久期代理"
    out.loc[out["source_type"].astype(str).str.contains("derived_style_bucket"), "proxy_type"] = "点时研究构造桶"
    out.loc[out["asset_code"].eq("rate_30y"), "proxy_type"] = "超长国债财富指数代理"
    out.loc[out["asset_code"].eq("alt_reits"), "proxy_type"] = "REITs 等权研究指数"
    out.loc[out["asset_code"].eq("alt_us_bond"), "proxy_type"] = "美债10Y曲线久期代理"
    return out.sort_values(["module", "asset_code"])


def profile_budgets() -> list[list[object]]:
    cfg = read_yaml("config/v1_strategy_config.yaml")
    rows: list[list[object]] = []
    for profile, values in cfg["product"]["profiles"].items():
        target = values["target_risk_contribution"]
        bounds = values["sleeve_bounds"]
        for sleeve in ["rates", "credit", "convertible_bonds", "equity", "alternatives"]:
            rows.append([profile, sleeve, float(target[sleeve]), float(bounds[sleeve][0]), float(bounds[sleeve][1]), "现金目标为0"])
    return rows


def main() -> None:
    assets = asset_data()
    summary = pd.read_csv(OUT / f"{VERSION}_three_profile_summary.csv")
    baselines = pd.read_csv(OUT / "v43_clean_calendar_control_variable_baselines.csv")
    bootstrap = pd.read_csv(OUT / f"{VERSION}_bootstrap_robustness.csv")
    portfolio_gate = json.loads((OUT / f"{VERSION}_portfolio_gate.json").read_text(encoding="utf-8"))
    model_gate = json.loads((OUT / f"{VERSION}_gate.json").read_text(encoding="utf-8"))
    metadata = json.loads((OUT / f"{VERSION}_metadata.json").read_text(encoding="utf-8"))

    workbook = Workbook()
    workbook.remove(workbook.active)

    overview = make_sheet(workbook, "总览", "固收+多资产动态配置：三层框架与数据底稿", "正式版本 v43 Clean Calendar；结果和数据口径自动来自当前 38 资产正式运行文件。")
    overview_rows = [
        ["正式版本", VERSION, "模型最后决策日", "2026-07-15"],
        ["资产池", "38资产、5袖套", "训练窗口", "120个交易日"],
        ["预测期限", "5 / 20 / 60交易日", "执行", "T+1，5bp x 单边换手，无现金"],
        ["A0", "分层风险平价", "A1", "机械 Market View + EP/skfolio"],
        ["A2", "iTransformer 因果相对观点 + PPO 择时执行", "正式门控", str(portfolio_gate["enabled_as_formal"])],
        ["当前结论", "中高波有更强正向证据；低波 Bootstrap Sharpe 稳健性不足，保留 A1 锚点", "训练/回测状态", "已完成"],
    ]
    write_table(overview, 4, ["项目", "口径", "项目", "口径"], overview_rows, "OverviewTable")
    overview.merge_cells("A12:H12")
    overview["A12"] = "三层结构：A0 提供风险分散底仓；A1 将状态与观点转为受约束目标组合；A2 在因果门控下仅提供局部增量观点，PPO 只决定是否调仓和调多少。"
    overview["A12"].fill = PatternFill("solid", fgColor=PALE_BLUE)
    overview["A12"].alignment = Alignment(wrap_text=True, vertical="center")
    overview.row_dimensions[12].height = 38
    format_widths(overview, {"A": 18, "B": 42, "C": 18, "D": 34, "E": 14, "F": 14, "G": 14, "H": 14})

    framework = make_sheet(workbook, "三层框架", "三层策略框架", "每层均以同一资产池、同一交易日、T+1 和 5bp 成本为比较前提。")
    framework_rows = [
        ["A0", "分层风险平价底仓", "解决长期风险分散与产品分层", "五袖套历史收益、协方差、风险贡献目标、上下限", "袖套间风险预算；袖套内逆波动", "月度目标权重、日频净值", "不使用观点、iTransformer 或 PPO"],
        ["A1", "Market View + EP/skfolio", "解决可解释的主动配置", "状态仪表盘、股债九宫格、38资产卡、机械观点、收缩协方差", "观点修正期望收益；风险预算和边界约束仓位", "受约束目标组合", "正式锚点，A2 的直接对照"],
        ["A2", "iTransformer + PPO", "解决局部相对机会与日频执行择时", "过去120日特征、状态卡、九宫格、机械观点、当前权重、成本", "iTransformer 预测相对观点；因果门控；PPO 选择持有或向目标调25/50/100%", "日频实际持仓、交易记录、净值", "模型不直接造仓位；PPO 不创造方向"],
    ]
    write_table(framework, 4, ["层级", "模块", "核心问题", "主要输入", "规则/算法", "主要输出", "边界"], framework_rows, "LayerTable")
    framework.merge_cells("A10:G10")
    framework["A10"] = "九宫格定位：债券强/中/弱 × 权益强/中/弱，是状态条件变量和跨资产背景，不直接替代信用、转债、另类和单资产判断。"
    framework["A10"].fill = PatternFill("solid", fgColor=PALE_GREEN)
    framework["A10"].alignment = Alignment(wrap_text=True, vertical="center")
    framework.row_dimensions[10].height = 34
    format_widths(framework, {"A": 10, "B": 24, "C": 26, "D": 43, "E": 42, "F": 25, "G": 28})

    assets_sheet = make_sheet(workbook, "38资产字典", "38资产主数据字典", "“代理类型”明确区分可交易价格、指数/曲线代理和点时研究构造桶；数据起止日来自当前收益面板。", "K")
    asset_rows = []
    for _, row in assets.iterrows():
        asset_rows.append([
            row["module"], row["asset_code"], row["name"], row.get("ts_code", ""), row["data_source"],
            row["proxy_type"], pd.Timestamp(row["actual_start"]).date() if pd.notna(row["actual_start"]) else None,
            pd.Timestamp(row["actual_end"]).date() if pd.notna(row["actual_end"]) else None,
            int(row["observations"]), float(row["missing_rate"]), row.get("note", ""),
        ])
    end_row, _ = write_table(assets_sheet, 4, ["袖套", "资产代码", "资产名称", "来源代码", "数据源", "代理类型", "实际开始", "实际结束", "观测数", "日收益缺失率", "说明"], asset_rows, "AssetDictionary")
    for row in range(5, end_row + 1):
        assets_sheet.cell(row, 7).number_format = "yyyy-mm-dd"
        assets_sheet.cell(row, 8).number_format = "yyyy-mm-dd"
        assets_sheet.cell(row, 10).number_format = "0.00%"
    assets_sheet.conditional_formatting.add(f"J5:J{end_row}", ColorScaleRule(start_type="min", start_color="E2F0D9", mid_type="percentile", mid_value=50, mid_color="FFF2CC", end_type="max", end_color="F4CCCC"))
    format_widths(assets_sheet, {"A": 18, "B": 30, "C": 24, "D": 18, "E": 31, "F": 25, "G": 13, "H": 13, "I": 10, "J": 14, "K": 60})

    data_sheet = make_sheet(workbook, "数据文件", "数据文件、频率与点时边界", "本表描述每个文件的责任边界，而非仅列出文件名。")
    data_rows = [
        ["资产收益与OHLCV", "data_processed/asset_returns_daily.parquet", "日频/交易日", "38资产收益、局部特征、回测", "仅用当日及以前记录；源切换日收益剔除"],
        ["日频特征面板", "data_processed/feature_panel_daily_v43.parquet", "日频/交易日", "利率、信用、资金、转债、权益、另类状态", "滚动计算只向后看"],
        ["月频宏观面板", "data_processed/macro_factors_monthly_v43.parquet", "月频", "宏观背景和低频状态", "按真实发布后可得日对齐"],
        ["市场状态卡", "data_processed/market_state_cards_daily_v1.parquet", "日频/交易日", "仪表盘、九宫格、机械观点、模型上下文", "只由截至当日特征生成"],
        ["机械Market View", "data_processed/market_skill_views_mechanical_daily_v1.parquet", "日频/交易日", "A1观点与模型输入", "规则在当日状态卡上生成"],
        ["模型数据集", "data_processed/hierarchical_asset_itransformer_v43_clean_calendar_dataset.npz", "120日滚动窗口", "模型输入与5/20/60日标签", "标签严格从t+1开始"],
        ["模型预测", f"data_processed/{VERSION}_asset_predictions.parquet", "日频/交易日", "资产与袖套相对观点", "不通过门控时严格回退为零"],
        ["最终回测", f"output/reports/{VERSION}_*_daily.csv", "日频/交易日", "PPO、机械、正式锚点的同日期回放", "T+1、5bp成本、无现金"],
    ]
    write_table(data_sheet, 4, ["数据类别", "文件", "频率", "主要用途", "点时/质量规则"], data_rows, "DataFiles")
    format_widths(data_sheet, {"A": 21, "B": 70, "C": 18, "D": 35, "E": 38})

    features = make_sheet(workbook, "特征与标签", "iTransformer 特征与标签口径", "特征为过去120个交易日的因果窗口；协方差风险不作为未来收益标签的一部分。")
    feature_rows = [
        ["局部资产特征", "38资产日收益、动量、波动、回撤及缺失掩码", "资产自身价格与风险轨迹", "每日", "不以未上市/缺失替代为零收益"],
        ["利率与信用", "国债曲线、信用利差、DR007、Shibor", "利率久期、信用和资金压力", "每日", "滚动变化只使用已见历史"],
        ["转债与权益", "转债价格/溢价中位数、CSI1000 PE/PB、权益广度", "估值、风险偏好与市场广度", "每日", "估值按可得日进入"],
        ["另类宏观", "美债10Y、美元、南华商品、黄金、实际利率、REITs分派率", "商品、避险与海外利率背景", "日频代理/低频补充", "缺失由掩码和点时前向规则处理"],
        ["状态与九宫格", "六类状态码、股债3x3 one-hot、机械五袖套观点", "市场环境与规则化上下文", "每日", "九宫格只做条件变量"],
        ["预测标签", "未来t+1至t+5/20/60日年化对数平均收益，横截面去均值", "相对表现而非绝对价格方向", "5/20/60日", "标签必须在决策日后成熟"],
    ]
    write_table(features, 4, ["特征组", "代表字段/构造", "经济含义", "更新", "防泄露规则"], feature_rows, "FeatureTable")
    features.merge_cells("A13:H13")
    features["A13"] = "当前元数据记录：局部特征数量 " + str(len(metadata.get("local_features", []))) + "；宏观/状态特征数量 " + str(len(metadata.get("macro_features", []))) + "。具体字段列表见数据底稿 Markdown，Excel 保留可读的分组口径。"
    features["A13"].fill = PatternFill("solid", fgColor=PALE_BLUE)
    features["A13"].alignment = Alignment(wrap_text=True, vertical="center")
    features.row_dimensions[13].height = 35
    format_widths(features, {"A": 20, "B": 52, "C": 28, "D": 18, "E": 38})

    budgets = make_sheet(workbook, "产品风险预算", "三产品风险预算与边界", "风险预算是A0的战略底仓输入；名义权重会随协方差与上下限变化。")
    budget_rows = profile_budgets()
    end_row, _ = write_table(budgets, 4, ["产品", "袖套", "目标风险贡献", "权重下限", "权重上限", "现金规则"], budget_rows, "RiskBudgets")
    for row in range(5, end_row + 1):
        for col in [3, 4, 5]:
            budgets.cell(row, col).number_format = "0.0%"
    format_widths(budgets, {"A": 16, "B": 22, "C": 18, "D": 14, "E": 14, "F": 20})

    results_sheet = make_sheet(workbook, "结果与门控", "A0、A1、A2 结果与正式门控", "所有结果为同资产池、同日期、T+1、5bp成本；不要将A2的局部改善表述为全产品正式替代。", "J")
    result_rows = []
    for _, row in baselines.iterrows():
        result_rows.append([row["profile"], row["strategy"], "2022-01-03 至 2026-07-15", row["annual_return"], row["annual_vol"], row["sharpe"], row["max_drawdown"], row["mean_daily_turnover"] * 252, row["total_cost"], "基线"])
    for _, row in summary.loc[summary["split"].eq("oos_2024_onward")].iterrows():
        result_rows.append([row["profile"], "A2 iTransformer + PPO", "2024-01-02 至 2026-07-15", row["strategy_annual_return"], row["strategy_annual_vol"], row["strategy_sharpe"], row["strategy_max_drawdown"], row["annualized_turnover"], None, f"相对A1 Sharpe {row['delta_sharpe_vs_formal']:+.3f}"])
    end_row, _ = write_table(results_sheet, 4, ["产品", "版本", "评价区间", "年化收益", "年化波动", "Sharpe", "最大回撤", "年化换手", "累计成本", "说明"], result_rows, "ResultsTable")
    for row in range(5, end_row + 1):
        for col in [4, 5, 7, 8, 9]:
            results_sheet.cell(row, col).number_format = "0.00%"
        results_sheet.cell(row, 6).number_format = "0.00"
    bootstrap_rows = []
    for _, row in bootstrap.iterrows():
        bootstrap_rows.append([row["profile"], int(row["bootstrap_draws"]), int(row["bootstrap_block_days"]), row["bootstrap_return_positive_probability"], row["bootstrap_sharpe_positive_probability"], row["bootstrap_return_edge_q05"], row["bootstrap_return_edge_q95"], row["bootstrap_sharpe_edge_q05"], row["bootstrap_sharpe_edge_q95"]])
    start = end_row + 3
    end_boot, _ = write_table(results_sheet, start, ["产品", "抽样次数", "块长度", "收益优势概率", "Sharpe优势概率", "收益优势P05", "收益优势P95", "Sharpe优势P05", "Sharpe优势P95"], bootstrap_rows, "BootstrapTable")
    for row in range(start + 1, end_boot + 1):
        for col in [4, 5, 6, 7]:
            results_sheet.cell(row, col).number_format = "0.00%"
        for col in [8, 9]:
            results_sheet.cell(row, col).number_format = "0.000"
    results_sheet.merge_cells(f"A{end_boot + 3}:J{end_boot + 3}")
    results_sheet.cell(end_boot + 3, 1).value = f"正式门控 = {portfolio_gate['enabled_as_formal']}。未通过原因：低波产品的 Bootstrap Sharpe 优势概率未达到预设阈值；中高波的证据更强。"
    results_sheet.cell(end_boot + 3, 1).fill = PatternFill("solid", fgColor=PALE_AMBER)
    results_sheet.cell(end_boot + 3, 1).alignment = Alignment(wrap_text=True, vertical="center")
    results_sheet.row_dimensions[end_boot + 3].height = 34
    format_widths(results_sheet, {"A": 14, "B": 29, "C": 24, "D": 14, "E": 14, "F": 12, "G": 14, "H": 14, "I": 14, "J": 26})

    limits = make_sheet(workbook, "代理与限制", "代理资产、数据限制与文章披露要求", "这一页是对外写作时必须保留的边界，不因回测表现较好而删除。")
    limit_rows = [
        ["30Y利率债", "10年以上国债财富指数作为30Y可回测久期代理", "不等同纯30Y指数或ETF；原始名rate_30y、模型规范名rate_30y_curve_proxy不可重复配置", "明确披露"],
        ["美债", "美国10Y收益率构造7.5久期代理", "非可直接交易ETF；汇率与QDII交易摩擦不完全等价", "明确披露"],
        ["REITs", "首批9只公募REITs日收益等权研究指数", "2021-06-22起可得，非官方可交易指数", "明确披露"],
        ["城投AA/AA+", "iFind细分指数增量数据", "评级细分口径与源切换需持续审计", "明确披露"],
        ["城投AAA", "城投债ETF代理", "不能表述为严格评级拆分的城投AAA指数", "明确披露"],
        ["转债风格桶", "前一交易日价格与余额点时分桶", "研究组合，未纳入完整成分级冲击、停牌与交易成本", "明确披露"],
        ["权益行业代理", "部分使用底层指数而非ETF", "不含ETF跟踪误差与申赎摩擦", "明确披露"],
        ["月频宏观", "按发布日点时对齐", "不能将未发布当月值提前写入特征或标签", "严格执行"],
        ["模型观点", "部分模块/年份/期限经因果门控放行", "不能声称模型每天对38资产全面给出有效预测", "严格执行"],
        ["PPO", "执行择时控制器", "不创造方向、不突破风险预算、不直接选资产", "严格执行"],
    ]
    write_table(limits, 4, ["对象", "当前处理", "限制/风险", "对外披露"], limit_rows, "LimitTable")
    format_widths(limits, {"A": 20, "B": 44, "C": 58, "D": 16})

    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is not None and cell.row > 2:
                    font = copy(cell.font)
                    font.name = "Microsoft YaHei"
                    cell.font = font
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.page_margins.left = 0.25
        sheet.page_margins.right = 0.25
        sheet.page_margins.top = 0.4
        sheet.page_margins.bottom = 0.4

    workbook.save(OUTPUT)

    # Lightweight integrity check after serialization.
    checked = load_workbook(OUTPUT, read_only=True, data_only=False)
    expected = ["总览", "三层框架", "38资产字典", "数据文件", "特征与标签", "产品风险预算", "结果与门控", "代理与限制"]
    if checked.sheetnames != expected:
        raise RuntimeError(f"Workbook sheet check failed: {checked.sheetnames}")
    if checked["38资产字典"].max_row != 42:
        raise RuntimeError("Asset dictionary must contain exactly 38 asset rows.")
    print(f"created: {OUTPUT}")


if __name__ == "__main__":
    main()
