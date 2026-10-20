"""Snapshot the authoritative local CB database and rebuild isolated research inputs.

Never writes to the external database, frozen V45 data or published portfolios.
The historical-balance variant is diagnostic, not a certified PIT backtest.
"""
from pathlib import Path
import argparse
from datetime import datetime, timezone
import hashlib
import itertools
import json
import shutil
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import numpy as np
import pandas as pd
from src.v49_cb_database import (panel, historical_universe, observed_returns,
                                rebalance_turnover, strict_nav)
from src.v49_point_in_time import (cb_disclosure_events, align_disclosed_balances,
                                  cb_bucket_weights, realized_bucket_return)
from src.v45_forecast import write_json

DEFAULT = Path.home()/"Desktop/可转债数据库生成"
FIELDS = {
    "open":"1_Base_Data/转债开盘价_tushare.parquet",
    "high":"1_Base_Data/转债最高价_tushare.parquet",
    "low":"1_Base_Data/转债最低价_tushare.parquet",
    "close":"1_Base_Data/转债收盘价_tushare.parquet",
    "preclose":"1_Base_Data/昨收盘价_tushare.parquet",
    "volume":"1_Base_Data/成交量_tushare.parquet",
    "amount":"1_Base_Data/成交额_tushare.parquet",
    "balance":"1_Base_Data/债券余额.parquet",
    "premium":"1_Base_Data/平价溢价率_tushare.parquet",
    "parity":"1_Base_Data/平价_tushare.parquet",
    "conversion_price":"1_Base_Data/转股价.parquet",
    "ytm":"1_Base_Data/YTM.parquet",
    "remaining_years":"1_Base_Data/剩余期限.parquet",
    "iv_strict":"2_Derived_Data/隐含波动率_bt.parquet",
    "iv_adjusted":"2_Derived_Data/隐含波动率_adj.parquet",
}
NAMES = [f"cb_{a}_{b}_{c}" for a,b,c in itertools.product(
    ["bond", "balanced", "equity"], ["large", "small"], ["low", "high"])]


def digest(path):
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda:f.read(1024*1024), b""):
            h.update(chunk)
    return h.hexdigest()


def snapshot(source, target, entries):
    before = digest(source)
    target.parent.mkdir(parents=True, exist_ok=True)
    if target.exists():
        if digest(target) != before:
            raise ValueError(f"Snapshot changed; use a new --tag: {target.name}")
    else:
        shutil.copy2(source, target)
    if digest(source) != before or digest(target) != before:
        raise ValueError("Source mutated during read-only snapshot")
    entries.append(dict(source=str(source), snapshot=str(target.relative_to(ROOT)),
                        sha256=before, bytes=target.stat().st_size))
    return target


def verify_excel(workbook, frames, reports):
    # One shared-string load; exact named sheets, never positional sheet IDs.
    from openpyxl import load_workbook
    w = load_workbook(workbook, read_only=True, data_only=True)
    checks = []
    try:
        for sheet, field in [("cb_转债收盘价", "close"), ("cb_债券余额", "balance")]:
            rows = w[sheet].iter_rows(values_only=True)
            headers = next(rows)
            reference = frames[field]
            # Streaming XLSX rows can omit trailing empty cells entirely.
            records = []
            for row in rows:
                if len(row)>len(headers):
                    raise ValueError(f"Unexpected worksheet cells beyond header: {sheet}")
                records.append(tuple(row)+(None,)*(len(headers)-len(row)))
            frame = panel(pd.DataFrame.from_records(records, columns=headers),
                          start=str(reference.index.min().date()),
                          cutoff=str(reference.index.max().date()))
            if set(frame.columns) != set(reference.columns) or not frame.index.equals(reference.index):
                raise ValueError(f"Workbook/source schema mismatch: {sheet}")
            delta = frame-reference.reindex(columns=frame.columns)
            missing_mismatch = frame.isna().ne(reference.reindex(columns=frame.columns).isna())
            max_error = float(delta.abs().max().max())
            passed = not missing_mismatch.any().any() and max_error < 1e-7
            checks.append(dict(sheet=sheet, rows=len(frame), columns=len(frame.columns),
                               max_absolute_difference=max_error, passed=passed))
            if not passed:
                raise ValueError(f"Workbook/source value mismatch: {sheet}")
            print(f"Verified workbook {sheet}: {len(frame)} dates", flush=True)
    finally:
        w.close()
    pd.DataFrame(checks).to_csv(reports/"workbook_source_parity.csv", index=False)


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--source-directory", type=Path, default=DEFAULT)
    ap.add_argument("--cutoff", default="2026-10-14")
    ap.add_argument("--tag", default="20261014_v49")
    args = ap.parse_args()
    if not args.tag.replace("_", "").isalnum():
        raise ValueError("Tag must be a single safe directory name")
    cutoff = pd.Timestamp(args.cutoff)
    source = args.source_directory.resolve()
    cfg = json.loads((source/"config.json").read_text(encoding="utf-8"))
    data = Path(cfg["data_root"])
    # A dated production artifact is required; *_latest and intermediate files
    # must not win simply because somebody copied them more recently.
    stamp = cutoff.strftime("%Y%m%d")
    workbook = source/f"output/{stamp}/转债数据库_{stamp}.xlsx"
    if not workbook.is_file():
        raise FileNotFoundError(workbook)
    raw = ROOT/"data_raw"/f"cb_database_{args.tag}"
    out = ROOT/"data_processed"/f"cb_rebuild_{args.tag}"
    reports = ROOT/"output/reports"/f"cb_rebuild_{args.tag}"
    out.mkdir(parents=True, exist_ok=True)
    reports.mkdir(parents=True, exist_ok=True)
    progress = reports/"progress.json"
    write_json(progress, dict(status="running", stage="snapshot", completed_days=0))
    manifest = []
    try:
        wb = snapshot(workbook, raw/workbook.name, manifest)
        frames = {}
        coverage = []
        for name, relative in FIELDS.items():
            path = data/relative
            if not path.exists():
                if name.startswith("iv_"):
                    continue
                raise FileNotFoundError(path)
            local = snapshot(path, raw/relative, manifest)
            frames[name] = panel(pd.read_parquet(local), cutoff=args.cutoff)
            f = frames[name]
            coverage.append(dict(field=name, start=f.index.min(), end=f.index.max(),
                dates=len(f), securities=len(f.columns), observed_cells=int(f.notna().sum().sum()),
                missing_fraction=float(f.isna().to_numpy().mean()),
                release_verified=False if name not in {"open","high","low","close","preclose","volume","amount"} else True,
                role="research_only_fallback_mixture" if name=="iv_adjusted" else "individual_observations"))
            print(f"Snapshot {name}: {f.shape} through {f.index.max().date()}", flush=True)
        if frames["close"].index.max() != cutoff:
            raise ValueError("Newest actual quote date differs from requested cutoff")
        pd.DataFrame(coverage).to_csv(reports/"field_coverage.csv", index=False)
        verify_excel(wb, frames, reports)
        basic_path = snapshot(data/"1_Base_Data/_meta/balance_basic_last.parquet", raw/"basic.parquet", manifest)
        basic = pd.read_parquet(basic_path)
        master = historical_universe(basic, frames["close"].columns, args.cutoff)
        master.to_parquet(out/"historical_security_master_audited.parquet")
        codes = master.index[master.included_research_universe].tolist()
        dates = frames["close"].index
        close = frames["close"].reindex(columns=codes)
        returns = observed_returns(close, frames["preclose"])
        # Listing is an immutable event, unlike current survivor/balance filters.
        for code in codes:
            before = dates < master.at[code,"list_date_parsed"]
            close.loc[before, code] = np.nan
            returns.loc[before, code] = np.nan
        returns.to_parquet(out/"individual_observed_price_returns.parquet")
        for field, f in frames.items():
            f.reindex(index=dates, columns=codes).to_parquet(out/f"{field}_individual.parquet")
        balance = frames["balance"].reindex(index=dates, columns=codes)*1e8
        issue = pd.DataFrame(np.broadcast_to(master.loc[codes,"issue_size"].to_numpy(), balance.shape), index=dates, columns=codes)
        events_paths = [data/"1_Base_Data/_meta/tushare_cb_share_event_codes.parquet",
            ROOT/"data_raw/refresh_20261014_wind_verified/supplement_v49/candidates/cb_share_dated_disclosures.parquet"]
        disclosures = []
        for i,p in enumerate(events_paths):
            copied = snapshot(p, raw/f"disclosures_{i}.parquet", manifest)
            f = pd.read_parquet(copied)[["ts_code","publish_date","end_date","remain_size","issue_size"]].copy()
            for col in ["publish_date","end_date"]:
                f[col] = pd.to_datetime(f[col], errors="raise")
            disclosures.append(f)
        events = cb_disclosure_events(pd.concat(disclosures, ignore_index=True), dates, args.cutoff)
        known = align_disclosed_balances(events, dates, codes, max_age_days=120)
        events.to_parquet(out/"dated_balance_announcements.parquet", index=False)
        known["balance"].to_parquet(out/"last_published_balance_yuan.parquet")
        known["age_days"].to_parquet(out/"balance_announcement_age_days.parquet")
        quoted = close.gt(0)
        features = pd.DataFrame(index=dates)
        features["cb_count_quoted_historical_universe"] = quoted.sum(axis=1)
        features["cb_price_median"] = close.where(quoted).median(axis=1)
        for field in ["premium","parity","ytm","remaining_years","iv_strict","iv_adjusted"]:
            if field in frames:
                f = frames[field].reindex(index=dates,columns=codes).where(quoted)
                features[f"cb_{field}_median"] = f.median(axis=1)
                features[f"cb_{field}_observed_count"] = f.notna().sum(axis=1)
        features["cb_amount_yuan"] = (frames["amount"].reindex(index=dates,columns=codes)*10000).where(quoted).sum(axis=1,min_count=1)
        features["cb_historical_balance_yuan_unverified_release"] = balance.where(quoted).sum(axis=1,min_count=1)
        features["cb_announced_balance_coverage"] = (known["balance"].gt(0)&quoted).sum(axis=1)/quoted.sum(axis=1).replace(0,np.nan)
        features.to_parquet(out/"cb_market_observed_features_no_interpolation.parquet")
        gap = close.gt(0) & balance.le(0).where(balance.notna(), True)
        rows, members, missing, previous_books = [], [], [], {}
        variants = {"historical_balance_diagnostic":(balance,issue),
                    "announced_balance_strict_candidate":(known["balance"],known["issue_size"])}
        for i in range(1, len(dates)):
            signal, day = dates[i-1], dates[i]
            for variant,(b,s) in variants.items():
                groups = cb_bucket_weights(close.loc[signal], b.loc[signal], s.loc[signal])
                for name in NAMES:
                    detail = groups.get(name)
                    prior = previous_books.get((variant,name))
                    if detail is None:
                        rows.append(dict(variant=variant,date=day,signal_date=signal,asset_code=name,
                            return_1d=np.nan, member_count=0,max_weight=np.nan,weight_sum=np.nan,
                            fallback_level=-1,one_way_turnover=np.nan,cost_5bp=np.nan))
                        previous_books.pop((variant,name), None)
                        continue
                    w = detail["weights"]
                    turnover = rebalance_turnover(prior, returns.loc[signal], w)
                    realized = realized_bucket_return(w,returns.loc[day])
                    rows.append(dict(variant=variant,date=day,signal_date=signal,asset_code=name,
                        return_1d=realized,member_count=len(w),max_weight=float(w.max()),weight_sum=float(w.sum()),
                        fallback_level=detail["fallback_level"],one_way_turnover=turnover,cost_5bp=turnover*.0005))
                    for code,weight in w.items():
                        members.append((variant,day,signal,name,code,float(weight)))
                        r = returns.at[day,code]
                        if not np.isfinite(r) or r<=-1:
                            missing.append(dict(variant=variant,date=day,signal_date=signal,asset_code=name,
                                ts_code=code,precommitted_weight=float(weight),prior_close=close.at[signal,code],
                                current_close=close.at[day,code],current_preclose=frames["preclose"].at[day,code],
                                resolution="unresolved_mark_or_terminal_cashflow_no_renormalization"))
                    previous_books[(variant,name)] = w
            if i%100 == 0 or i==len(dates)-1:
                print(f"Rebuilt {i}/{len(dates)-1} days, {day.date()}, missing positions={len(missing)}", flush=True)
                write_json(progress,dict(status="running",stage="buckets",completed_days=i,total_days=len(dates)-1))
        result = pd.DataFrame(rows)
        result["net_return_5bp_diagnostic"] = result.return_1d-result.cost_5bp
        result["nav_no_gap_fill"] = result.groupby(["variant","asset_code"]).return_1d.transform(strict_nav)
        result.to_parquet(out/"cb_12_buckets_rebuilt.parquet", index=False)
        pd.DataFrame(members,columns=["variant","date","signal_date","asset_code","ts_code","weight"]).to_parquet(out/"precommitted_holdings.parquet",index=False)
        pd.DataFrame(missing).to_csv(reports/"unresolved_held_positions.csv",index=False)
        summary = result.groupby(["variant","asset_code"]).agg(days=("date","size"),
            valid_return_days=("return_1d","count"),average_members=("member_count","mean"),
            max_single_weight=("max_weight","max"),mean_one_way_turnover=("one_way_turnover","mean"),
            average_daily_cost_5bp=("cost_5bp","mean")).reset_index()
        summary.to_csv(reports/"bucket_coverage_and_turnover.csv",index=False)
        old_path = ROOT/"data_processed/asset_returns_daily.parquet"
        if old_path.exists():
            old = pd.read_parquet(old_path,columns=["date","asset_code","return_1d"])
            old["date"] = pd.to_datetime(old.date)
            if old.duplicated(["date","asset_code"]).any():
                raise ValueError("Legacy asset return keys are duplicated")
            comparison = result.merge(old,on=["date","asset_code"],how="left",suffixes=("_new","_old"))
            comparison["difference"] = comparison.return_1d_new-comparison.return_1d_old
            comparison.to_parquet(out/"legacy_return_comparison.parquet",index=False)
            comparison.groupby(["variant","asset_code"]).difference.agg(["count","mean","std","min","max"]).to_csv(reports/"legacy_return_differences.csv")
        write_json(raw/"manifest.json",dict(created_at=datetime.now(timezone.utc).isoformat(),
            cutoff=args.cutoff,files=manifest,config_secrets_copied=False))
        status = dict(status="rebuilt_candidates_not_promoted",cutoff=args.cutoff,
            first_date=str(dates.min().date()),dates=len(dates),source_security_columns=len(master),
            historical_candidate_codes=len(codes),codes_with_announced_balances=events.ts_code.nunique(),
            historical_quote_balance_missing_cells=int(gap.sum().sum()),
            unresolved_position_rows=len(missing),
            missing_bucket_returns=result.groupby("variant").return_1d.apply(lambda s:int(s.isna().sum())).to_dict(),
            actual_balance_vs_announced_balance_are_distinct=True,training_ready=False,
            blockers=["historical_issuance_method_registry_unverified",
                "historical_daily_balances_not_certified_publication_time",
                "dated_disclosure_coverage_incomplete_especially_before_2020",
                "missing_held_marks_terminal_payoffs_and_coupons_require_resolution",
                "daily_reconstitution_tradeability_and_internal_costs_not_yet_certified"],
            workbook=str(workbook),outputs=str(out),reports=str(reports),
            formal_files_overwritten=False,models_retrained=False)
        write_json(reports/"rebuild_status.json", status)
        write_json(progress,dict(status="completed_candidates_blocked",completed_days=len(dates)-1,total_days=len(dates)-1))
        print(json.dumps(status,ensure_ascii=False,indent=2),flush=True)
    except Exception as exc:
        write_json(progress,dict(status="failed",error=f"{type(exc).__name__}: {exc}"))
        raise


if __name__ == "__main__":
    main()
